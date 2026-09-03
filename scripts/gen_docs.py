# -*- coding: utf-8 -*-
"""Generates the Vallentin Claims business and technical documentation as .docx files."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- Brand palette ---------------------------------------------------------
RED = RGBColor(0xB4, 0x32, 0x32)
RED_DARK = RGBColor(0x7D, 0x1F, 0x1F)
INK = RGBColor(0x0C, 0x14, 0x16)
NAVY = RGBColor(0x00, 0x4C, 0xA0)
GREY = RGBColor(0x3A, 0x61, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

OUT_DIR = r"C:\private\Vallentin_Claims"


# --- Low-level helpers -----------------------------------------------------
def shade(el, fill):
    pr = el.get_or_add_tcPr() if el.tag.endswith('}tc') else el.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pr.append(shd)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_base_style(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15


def recolor_heading(paragraph, color):
    for run in paragraph.runs:
        run.font.color.rgb = color
        run.font.name = 'Calibri'


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    recolor_heading(p, RED)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    recolor_heading(p, RED_DARK)
    return p


def h3(doc, text):
    p = doc.add_heading(text, level=3)
    recolor_heading(p, NAVY)
    return p


def para(doc, text, italic=False, size=10.5, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    r.font.size = Pt(size)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.add_run(text)
    return p


def kv_bullet(doc, key, value):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(key + ' — ')
    r.bold = True
    p.add_run(value)
    return p


def mono_block(doc, text, fill='F2F0EC'):
    """Render a monospace, shaded single-cell table (for code / diagrams)."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    shade_cell(cell, fill)
    cell.paragraphs[0].text = ''
    for i, line in enumerate(text.split('\n')):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = INK
    _set_table_borders(tbl, 'D8D2C8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def _set_table_borders(tbl, color='BBBBBB'):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color)
        borders.append(e)
    tblPr.append(borders)


def table(doc, headers, rows, widths=None, header_fill='B43232'):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0].cells
    for i, htext in enumerate(headers):
        shade_cell(hdr[i], header_fill)
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(htext)
        run.bold = True
        run.font.color.rgb = WHITE
        run.font.size = Pt(9.5)
    for ridx, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            if ridx % 2 == 1:
                shade_cell(cells[i], 'F7F4F0')
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(str(val))
            run.font.size = Pt(9)
    _set_table_borders(tbl)
    if widths:
        for i, w in enumerate(widths):
            for row in tbl.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def add_field(paragraph, field):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = field
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "Right-click here and choose 'Update Field' to generate the table of contents."
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    r = run._r
    for x in (f1, it, f2, t, f3):
        r.append(x)


def setup_footer(doc, title):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ''
    r = p.add_run(title + '    |    Page ')
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    add_field(p, 'PAGE')
    r2 = p.add_run(' of ')
    r2.font.size = Pt(8); r2.font.color.rgb = GREY
    add_field(p, 'NUMPAGES')


def cover(doc, kicker, title, subtitle, meta_rows):
    # top brand bar
    bar = doc.add_paragraph()
    br = bar.add_run('VALLENTIN CLAIMS')
    br.bold = True
    br.font.size = Pt(13)
    br.font.color.rgb = RED
    br.font.name = 'Calibri'
    doc.add_paragraph().paragraph_format.space_after = Pt(60)

    kp = doc.add_paragraph()
    kr = kp.add_run(kicker.upper())
    kr.bold = True
    kr.font.size = Pt(11)
    kr.font.color.rgb = GREY

    tp = doc.add_paragraph()
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(30)
    tr.font.color.rgb = INK
    tp.paragraph_format.space_after = Pt(4)

    sp = doc.add_paragraph()
    sr = sp.add_run(subtitle)
    sr.font.size = Pt(13)
    sr.font.color.rgb = GREY
    sp.paragraph_format.space_after = Pt(80)

    tbl = doc.add_table(rows=0, cols=2)
    for k, v in meta_rows:
        cells = tbl.add_row().cells
        rk = cells[0].paragraphs[0].add_run(k)
        rk.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = RED_DARK
        rv = cells[1].paragraphs[0].add_run(v)
        rv.font.size = Pt(10)
        cells[0].width = Inches(2.0)
        cells[1].width = Inches(4.5)
    doc.add_page_break()


def toc_page(doc):
    h1(doc, 'Table of Contents')
    add_toc(doc)
    doc.add_page_break()


# ===========================================================================
# BUSINESS DOCUMENT
# ===========================================================================
def build_business():
    doc = Document()
    set_base_style(doc)
    setup_footer(doc, 'Vallentin Claims — Business Documentation')
    doc.core_properties.author = 'Vallentin Claims Team'
    doc.core_properties.title = 'Vallentin Claims — Business Documentation'
    doc.core_properties.category = 'Business Documentation'

    cover(doc, 'Business Documentation',
          'Vallentin Claims',
          'Healthcare Complaint Management Platform',
          [('Document', 'Business Requirements & Product Overview'),
           ('Version', '1.0'),
           ('Status', 'Baseline'),
           ('Date', 'June 2026'),
           ('Owner', 'Product / Business Owner'),
           ('Classification', 'Internal')])

    toc_page(doc)

    # 1. Document control
    h1(doc, '1. Document Control')
    table(doc, ['Version', 'Date', 'Author', 'Summary of change'],
          [['0.1', 'May 2026', 'Product Owner', 'Initial draft of scope and objectives'],
           ['1.0', 'June 2026', 'Product Owner', 'Baseline aligned with delivered platform']],
          widths=[0.8, 1.1, 1.7, 3.0])
    para(doc, 'This document describes the business context, objectives, stakeholders, '
              'processes, and rules of the Vallentin Claims platform. It is the reference '
              'for product decisions and is written for a non-technical audience. The '
              'companion Technical Documentation covers architecture and implementation.')

    # 2. Executive summary
    h1(doc, '2. Executive Summary')
    para(doc, 'Vallentin Claims is a web platform that lets citizens submit complaints '
              'against healthcare providers — hospitals, individual doctors, and health '
              'insurance funds — and have them routed automatically to the responsible '
              'institution, always with a copy to the Ombudsman. The platform removes the '
              'friction of knowing where and how to complain: the citizen fills in a single '
              'structured form, and the system determines the recipient, records the '
              'submission, and forwards it by email with a full audit trail.')
    para(doc, 'The platform supports both registered users, who can track the complaints '
              'they have filed, and anonymous submitters, whose complaints are reviewed by '
              'an administrator before being forwarded. Every complaint receives a unique '
              'public reference number. Administrators moderate the anonymous queue, manage '
              'the register of institutions, oversee users, and export statistics for '
              'reporting. The product is fully bilingual (English and Bulgarian) and is '
              'designed to align visually with the Thirst-for-Life brand.')

    # 3. Business context
    h1(doc, '3. Business Context & Problem Statement')
    h2(doc, '3.1 The problem')
    para(doc, 'Patients and citizens frequently have legitimate grievances about healthcare '
              'services but face practical barriers to raising them:')
    bullet(doc, 'They do not know which institution is responsible or how to reach it.')
    bullet(doc, 'Complaints submitted informally leave no record and are easily lost.')
    bullet(doc, 'There is no consistent copy to an oversight body such as the Ombudsman.')
    bullet(doc, 'Vulnerable people may fear repercussions and need an anonymous channel.')
    h2(doc, '3.2 The solution')
    para(doc, 'Vallentin Claims provides a single, trustworthy intake point. It standardises '
              'the complaint, routes it to the correct institution, guarantees an oversight '
              'copy, issues a tracking reference, and preserves an immutable record of every '
              'step. Anonymous complaints are supported but pass through human moderation to '
              'protect institutions from abuse.')

    # 4. Vision & objectives
    h1(doc, '4. Vision & Objectives')
    kv_bullet(doc, 'Vision', 'Make it effortless and safe for any citizen to hold healthcare '
              'providers accountable, while giving oversight bodies reliable visibility.')
    para(doc, 'Business objectives:', bold=True)
    numbered(doc, 'Lower the barrier to filing a healthcare complaint to a single online form.')
    numbered(doc, 'Guarantee correct routing to the responsible institution and the Ombudsman.')
    numbered(doc, 'Provide transparency to submitters through a public tracking reference.')
    numbered(doc, 'Protect the process with moderation of anonymous complaints.')
    numbered(doc, 'Give administrators the tools to manage institutions, users, and reporting.')
    numbered(doc, 'Comply with data-protection obligations (consent, export, anonymisation).')

    # 5. Stakeholders
    h1(doc, '5. Stakeholders')
    table(doc, ['Stakeholder', 'Interest in the platform'],
          [['Citizen / complainant', 'Wants a simple, credible way to complain and to track progress.'],
           ['Anonymous complainant', 'Needs to raise a concern without revealing identity.'],
           ['Administrator / moderator', 'Reviews anonymous complaints, manages institutions and users.'],
           ['Ombudsman', 'Receives a copy of every forwarded complaint for oversight.'],
           ['Healthcare institution', 'Receives complaints addressed to it and is expected to respond.'],
           ['Business / product owner', 'Owns scope, brand alignment, and success metrics.'],
           ['Data protection officer', 'Ensures consent, retention, and subject-rights handling.']],
          widths=[1.9, 4.6])

    # 6. User roles
    h1(doc, '6. User Roles')
    table(doc, ['Role', 'Capabilities'],
          [['Anonymous user', 'Submit a complaint (subject to admin moderation before forwarding). '
            'No login required. Provides optional contact details.'],
           ['Authorised (registered) user', 'Register, verify email or sign in with Google/Facebook, '
            'submit complaints that are forwarded immediately, and view their own submissions.'],
           ['Administrator', 'Full management: moderate the anonymous queue (approve/reject), '
            'manage institutions, manage users, run and export statistics, and view audit history.']],
          widths=[1.9, 4.6])

    # 7. Personas
    h1(doc, '7. User Personas')
    h3(doc, 'Maria — the concerned patient')
    para(doc, 'Maria had a poor experience at a hospital and wants it recorded formally. She '
              'registers, verifies her email, submits a detailed complaint with a scanned '
              'document, and later checks its status. She values a clear reference number and '
              'reassurance that the Ombudsman was copied.')
    h3(doc, 'Georgi — the cautious whistle-blower')
    para(doc, 'Georgi works near the sector and fears repercussions. He submits anonymously, '
              'optionally leaving a contact email. His complaint waits in the moderation queue '
              'until an administrator approves it.')
    h3(doc, 'Elena — the administrator')
    para(doc, 'Elena reviews the anonymous queue daily, approves genuine complaints, rejects '
              'spam with a reason, keeps the institution register current, and exports monthly '
              'statistics for management reporting.')

    # 8. Business processes
    h1(doc, '8. Core Business Processes')
    h2(doc, '8.1 Registered complaint (happy path)')
    numbered(doc, 'The user signs in (email/password or social login).')
    numbered(doc, 'The user completes the complaint form: category, institution (from the '
             'register or free text), title, description, urgency, and optional attachments.')
    numbered(doc, 'The system validates the submission and creates the complaint with a public '
             'reference (format VLC-YYYY-NNNNNN).')
    numbered(doc, 'The complaint is forwarded automatically: an email to the institution, a '
             'mandatory copy to the Ombudsman, and an optional copy to the submitter.')
    numbered(doc, 'The user can view the complaint and its status among their submissions.')

    h2(doc, '8.2 Anonymous complaint (moderated path)')
    numbered(doc, 'An anonymous visitor completes the same form and passes a captcha check.')
    numbered(doc, 'The complaint is created with status "pending review" and is NOT forwarded yet.')
    numbered(doc, 'An administrator is notified and reviews the complaint in the moderation queue.')
    numbered(doc, 'On approval, the complaint is forwarded exactly like a registered one.')
    numbered(doc, 'On rejection, the administrator records a reason and nothing is forwarded.')

    h2(doc, '8.3 Complaint lifecycle (statuses)')
    table(doc, ['Status', 'Meaning'],
          [['Submitted', 'Registered-user complaint accepted; about to be forwarded.'],
           ['Pending review', 'Anonymous complaint awaiting administrator moderation.'],
           ['Approved', 'Anonymous complaint accepted by an administrator (then forwarded).'],
           ['Rejected', 'Anonymous complaint declined with a recorded reason; not forwarded.'],
           ['Forwarded', 'Emails dispatched to institution and Ombudsman.'],
           ['Closed', 'Complaint marked complete by an administrator.']],
          widths=[1.6, 4.9])

    # 9. Functional scope
    h1(doc, '9. Functional Scope (Feature Catalogue)')
    h2(doc, '9.1 Accounts & access')
    for t in ['Email/password registration with mandatory email verification.',
              'Social sign-in with Google and Facebook.',
              'Password reset by email; per-account lockout after repeated failed logins.',
              'Consent capture for Terms, Privacy, and optional Marketing, with versioning.']:
        bullet(doc, t)
    h2(doc, '9.2 Complaint handling')
    for t in ['Structured complaint form with category, institution, title, description, urgency.',
              'File attachments (PDF and common image formats) with size limits.',
              'Automatic routing to institution plus mandatory Ombudsman copy.',
              'Optional acknowledgement copy to the submitter.',
              'Unique public tracking reference per complaint.',
              'Moderation queue for anonymous complaints (approve / reject with reason).']:
        bullet(doc, t)
    h2(doc, '9.3 Administration')
    for t in ['Institution register: create, edit, activate/deactivate, delete; grouped by category.',
              'User management: list, change role/status, GDPR anonymisation, per-user data export.',
              'Complaint search and filtering by text, category, institution, status, and date.',
              'Full event history per complaint (immutable timeline).']:
        bullet(doc, t)
    h2(doc, '9.4 Statistics & reporting')
    for t in ['Public overview statistics on the home page (totals, forwarded, urgent, by category).',
              'Administrator statistics with date-range filtering.',
              'Export of complaint data to CSV and PDF.']:
        bullet(doc, t)
    h2(doc, '9.5 Platform')
    for t in ['Bilingual interface (English and Bulgarian).',
              'Brand-aligned visual design (Thirst-for-Life).',
              'Email delivery with automatic retry and dead-letter handling.']:
        bullet(doc, t)

    # 10. Business rules
    h1(doc, '10. Business Rules')
    table(doc, ['#', 'Rule'],
          [['BR-1', 'Every forwarded complaint is copied to the Ombudsman without exception.'],
           ['BR-2', 'Anonymous complaints must be approved by an administrator before forwarding.'],
           ['BR-3', 'Registered users\' complaints are forwarded automatically on submission.'],
           ['BR-4', 'A complaint description must meet a minimum length to be accepted.'],
           ['BR-5', 'Only complaints in "pending review" can be approved or rejected.'],
           ['BR-6', 'Only complaints in "forwarded" can be closed.'],
           ['BR-7', 'A complaint is forwarded at most once (no duplicate dispatch).'],
           ['BR-8', 'Attachments must match an allowed file type by actual content, not just name.'],
           ['BR-9', 'A user can only view complaints they submitted; administrators can view all.'],
           ['BR-10', 'Rejected complaints must carry a reason recorded by the administrator.']],
          widths=[0.7, 5.8])

    # 11. Compliance
    h1(doc, '11. Compliance & Data Protection')
    para(doc, 'The platform handles personal data and is designed with data-protection '
              'obligations in mind:')
    kv_bullet(doc, 'Consent', 'Terms and Privacy acceptance is captured at registration with '
              'the document version, timestamp, and originating IP/device recorded.')
    kv_bullet(doc, 'Right to export', 'Administrators can export all data held about a user.')
    kv_bullet(doc, 'Right to be forgotten', 'Administrators can anonymise a user, severing '
              'personal identifiers while preserving statistical integrity.')
    kv_bullet(doc, 'Data minimisation', 'Anonymous complaints require no account; contact '
              'details are optional.')
    kv_bullet(doc, 'Auditability', 'Security-relevant actions are recorded in an audit log.')

    # 12. Non-functional expectations
    h1(doc, '12. Non-Functional Expectations')
    table(doc, ['Area', 'Expectation'],
          [['Security', 'Strong password hashing, protected sessions, captcha, rate limiting, '
            'and abuse lockout.'],
           ['Privacy', 'Consent tracking, data export and anonymisation, minimal data capture.'],
           ['Reliability', 'Email delivery retried automatically; no complaint silently lost.'],
           ['Usability', 'Single-form submission, clear tracking reference, bilingual UI.'],
           ['Accessibility', 'Keyboard-navigable, labelled forms, readable contrast.'],
           ['Localisation', 'Full English and Bulgarian language support.'],
           ['Brand', 'Visual alignment with the Thirst-for-Life identity.']],
          widths=[1.6, 4.9])

    # 13. Success metrics
    h1(doc, '13. Success Metrics (KPIs)')
    table(doc, ['Metric', 'Why it matters'],
          [['Complaints submitted / month', 'Adoption and reach of the platform.'],
           ['% forwarded successfully', 'Core promise: complaints reach institutions.'],
           ['Median moderation time', 'Responsiveness for anonymous complaints.'],
           ['Rejection rate', 'Quality of intake and effectiveness of moderation.'],
           ['Email delivery success rate', 'Reliability of the routing guarantee.'],
           ['Registered vs anonymous ratio', 'Trust and engagement of the user base.']],
          widths=[2.3, 4.2])

    # 14. Risks
    h1(doc, '14. Risks & Assumptions')
    table(doc, ['Type', 'Item', 'Mitigation / note'],
          [['Risk', 'Abuse via anonymous spam', 'Captcha, rate limiting, mandatory moderation.'],
           ['Risk', 'Incorrect institution routing', 'Curated institution register; free-text fallback.'],
           ['Risk', 'Email non-delivery', 'Retry with dead-letter and audit visibility.'],
           ['Risk', 'Personal-data exposure', 'Access control, consent, anonymisation, audit.'],
           ['Assumption', 'Ombudsman email is stable', 'Configured centrally; single point to maintain.'],
           ['Assumption', 'Institutions act on emails', 'Platform delivers; response is out of scope.']],
          widths=[1.0, 2.6, 2.9])

    # 15. Roadmap
    h1(doc, '15. Out of Scope & Future Enhancements')
    para(doc, 'The current baseline delivers intake, routing, moderation, administration, and '
              'reporting. The following are candidates for future iterations:')
    bullet(doc, 'A truly public status lookup by reference number (no login required).')
    bullet(doc, 'Two-way correspondence between complainant and institution.')
    bullet(doc, 'Automated case escalation and SLA timers.')
    bullet(doc, 'Institution self-service portal for responding to complaints.')
    bullet(doc, 'Analytics dashboards and trend reporting.')
    bullet(doc, 'Embedding within the Thirst-for-Life WordPress site as a subdomain.')

    # 16. Glossary
    h1(doc, '16. Glossary')
    table(doc, ['Term', 'Definition'],
          [['Complaint', 'A structured grievance submitted against a healthcare institution.'],
           ['Public reference', 'The human-readable complaint identifier (VLC-YYYY-NNNNNN).'],
           ['Ombudsman', 'Oversight body that receives a copy of every forwarded complaint.'],
           ['Institution', 'A hospital, doctor, or health fund that a complaint targets.'],
           ['Moderation', 'Administrator review of anonymous complaints before forwarding.'],
           ['Forwarding', 'Automated email dispatch to the institution and Ombudsman.'],
           ['Anonymisation', 'Removal of personal identifiers from a user record.']],
          widths=[1.6, 4.9])

    path = os.path.join(OUT_DIR, 'Vallentin_Claims_Business_Documentation.docx')
    doc.save(path)
    return path


# ===========================================================================
# TECHNICAL DOCUMENT
# ===========================================================================
def build_technical():
    doc = Document()
    set_base_style(doc)
    setup_footer(doc, 'Vallentin Claims — Technical Documentation')
    doc.core_properties.author = 'Vallentin Claims Team'
    doc.core_properties.title = 'Vallentin Claims — Technical Documentation'
    doc.core_properties.category = 'Technical Documentation'

    cover(doc, 'Technical Documentation',
          'Vallentin Claims',
          'Architecture, Backend, Frontend & Communication Flows',
          [('Document', 'Technical Reference & Architecture'),
           ('Version', '1.0'),
           ('Status', 'Baseline'),
           ('Date', 'June 2026'),
           ('Owner', 'Engineering Lead'),
           ('Audience', 'Developers, DevOps, Architects')])

    toc_page(doc)

    # 1. Introduction
    h1(doc, '1. Introduction')
    para(doc, 'This document is the technical reference for the Vallentin Claims platform: a '
              'healthcare complaint management system consisting of a React single-page '
              'application, a Node.js/Express REST API, a PostgreSQL database, and an '
              'asynchronous email subsystem. It covers the architecture, module structure, '
              'data model, API surface, security design, communication flows with sequence '
              'diagrams, configuration, deployment, and operations.')

    # 2. Technology stack
    h1(doc, '2. Technology Stack')
    table(doc, ['Layer', 'Technology'],
          [['Frontend', 'React 18, Vite, TypeScript, TailwindCSS, React Router, TanStack Query, '
            'Zustand, react-i18next, Recharts'],
           ['Backend', 'Node.js 20, Express, TypeScript, Prisma ORM, Zod, Passport '
            '(Google + Facebook OAuth), Nodemailer, Pino'],
           ['Database', 'PostgreSQL 16 (with a GIN full-text index on complaints)'],
           ['Auth', 'JWT access + refresh tokens, Argon2id hashing, hCaptcha'],
           ['Email', 'Nodemailer over SMTP; MailHog in development'],
           ['Security', 'Helmet (CSP), CORS, express-rate-limit, file-type content sniffing'],
           ['Tooling', 'Jest + Supertest (backend), Playwright (E2E), ESLint, Prettier'],
           ['Deployment', 'Docker Compose (frontend, backend, postgres, mailhog)']],
          widths=[1.4, 5.1])

    # 3. High-level architecture
    h1(doc, '3. High-Level Architecture')
    para(doc, 'The system is a classic three-tier SPA + API + database architecture with an '
              'in-process background worker for asynchronous email delivery. All external '
              'communication with the API is over HTTPS/JSON; the browser holds a short-lived '
              'access token in memory and a long-lived refresh token in an HttpOnly cookie.')
    diagram = r"""
 +-----------------------------------------------------------------------+
 |                              Browser (SPA)                            |
 |   React 18 + Vite  |  Zustand (auth)  |  TanStack Query  |  i18n(EN/BG)|
 +--------------------------------+--------------------------------------+
                                  |  HTTPS / JSON  (Bearer access token,
                                  |                 refresh cookie)
                                  v
 +-----------------------------------------------------------------------+
 |                       Express API  (Node.js 20)                       |
 |  Middleware: requestId > logging > helmet(CSP) > CORS > cookies >     |
 |              json > passport > rate-limit                             |
 |                                                                       |
 |  Modules: auth | users | categories | institutions | complaints |    |
 |           attachments | statistics | admin | email | ombudsman |     |
 |           consents | auditLog                                        |
 |                                                                       |
 |  Background: Email Outbox Worker (setInterval, pg advisory lock)      |
 +-------------------+-----------------------------+---------------------+
                     |                             |
        Prisma ORM   |                             | Nodemailer (SMTP)
                     v                             v
 +-----------------------------+       +-------------------------------+
 |       PostgreSQL 16         |       |   SMTP server / MailHog       |
 |  users, complaints, events, |       |   -> Institution mailbox      |
 |  attachments, email_outbox, |       |   -> Ombudsman mailbox        |
 |  audit_log, ...             |       |   -> Submitter (optional)     |
 +-----------------------------+       +-------------------------------+
"""
    mono_block(doc, diagram.strip('\n'))

    # 4. Backend architecture
    h1(doc, '4. Backend Architecture')
    h2(doc, '4.1 Module layout')
    para(doc, 'The backend is organised into feature modules under backend/src/modules. Each '
              'module typically exposes a router, controller, service (business logic), and '
              'schema (Zod validation).')
    table(doc, ['Module', 'Responsibility'],
          [['auth', 'Registration, email verification, login, JWT refresh rotation, logout, '
            'password reset, Google/Facebook OAuth, one-time code exchange.'],
           ['users', 'Current-user profile, admin user management, anonymisation, data export.'],
           ['categories', 'Read the fixed set of complaint categories.'],
           ['institutions', 'Public listing/search; admin CRUD of the institution register.'],
           ['complaints', 'Submission, listing, detail, moderation, status changes, search, events.'],
           ['attachments', 'Access-controlled streaming of uploaded files.'],
           ['statistics', 'Public summary, admin detail, CSV/PDF export.'],
           ['admin', 'Aggregates admin sub-routers under /api/admin.'],
           ['email', 'Outbox enqueue, repository, and the background delivery worker.'],
           ['ombudsman', 'Builds the mandatory Ombudsman copy of each forwarded complaint.'],
           ['consents', 'Persists Terms/Privacy/Marketing consent records.'],
           ['auditLog', 'Writes security and lifecycle audit entries.']],
          widths=[1.4, 5.1])

    h2(doc, '4.2 Cross-cutting middleware')
    para(doc, 'Requests pass through a fixed middleware chain (see app.ts):')
    diagram2 = ("request  ->  requestId  ->  pino-http logging  ->  helmet (CSP)  ->\n"
                "CORS (credentialed)  ->  cookie-parser  ->  express.json (1mb)  ->\n"
                "passport.initialize  ->  global rate-limit  ->  route handlers  ->\n"
                "errorHandler (uniform JSON error envelope)")
    mono_block(doc, diagram2)
    para(doc, 'Additional per-route middleware includes requireAuth / optionalAuth (JWT), '
              'requireRole (RBAC), validate (Zod), requireCaptcha (hCaptcha), and route-specific '
              'rate limiters for auth and complaint submission.')

    h2(doc, '4.3 Error handling & configuration')
    para(doc, 'All handlers throw a typed HttpError (status, code, message, optional details) '
              'that the central errorHandler serialises into a consistent JSON envelope. '
              'Configuration is validated at startup with Zod (config/index.ts); the process '
              'refuses to start in production if placeholder secrets are detected.')

    # 5. Frontend architecture
    h1(doc, '5. Frontend Architecture')
    para(doc, 'The frontend is a Vite-built React SPA in TypeScript. Server state is managed by '
              'TanStack Query, authentication state by Zustand, routing by React Router, and '
              'internationalisation by react-i18next (English and Bulgarian).')
    h2(doc, '5.1 Page structure')
    table(doc, ['Area', 'Pages'],
          [['Public', 'Home (stats + tracking), About, Privacy, Terms, Track, NotFound, Forbidden'],
           ['Auth', 'Login, Register, ForgotPassword, ResetPassword, VerifyEmail, OAuthCallback'],
           ['Complaints', 'Submit, MyComplaints, ComplaintDetail'],
           ['Admin', 'Dashboard, ComplaintsQueue, Institutions, Users, Statistics']],
          widths=[1.4, 5.1])
    h2(doc, '5.2 Client-side concerns')
    bullet(doc, 'A shared API client attaches the access token and transparently refreshes it.')
    bullet(doc, 'Design tokens (brand red #b43232, navy, accent, ink) live in tailwind.config.ts.')
    bullet(doc, 'Reusable UI primitives (Button, Card, Badge, Spinner) enforce brand styling.')
    bullet(doc, 'Public runtime config is fetched from /api/config/public (limits, captcha, providers, locales).')

    # 6. Data model
    h1(doc, '6. Data Model')
    para(doc, 'The schema is defined in Prisma (backend/prisma/schema.prisma) and targets '
              'PostgreSQL. Key entities and relationships:')
    table(doc, ['Entity', 'Purpose & key fields'],
          [['User', 'Account: email, passwordHash, role (admin/user), status, lockout fields, '
            'lastLoginAt, anonymizedAt.'],
           ['OAuthIdentity', 'Links a user to a Google/Facebook identity (provider + providerUserId).'],
           ['Consent', 'Terms/Privacy/Marketing acceptance with version, timestamp, IP, UA.'],
           ['RefreshToken', 'Hashed refresh tokens with issue/expiry/revocation for rotation.'],
           ['EmailVerificationToken / PasswordResetToken', 'Hashed, single-use, expiring tokens.'],
           ['Category', 'Fixed complaint categories; parent of institutions and complaints.'],
           ['Institution', 'Target register: name, email, active flag, category.'],
           ['Complaint', 'Core entity: publicId, category, institution or free text, title, body, '
            'urgent, status, submissionType, reviewer, timestamps, IP/UA.'],
           ['Attachment', 'Uploaded file metadata linked to a complaint.'],
           ['ComplaintEvent', 'Immutable per-complaint timeline (created, forwarded, etc.).'],
           ['EmailOutbox', 'Queued email with status, attempts, next-attempt, error.'],
           ['AuditLog', 'Security/lifecycle audit trail keyed by actor and event.'],
           ['Configuration', 'Key/value runtime configuration.'],
           ['OAuthExchangeCode', 'One-time short-lived code to exchange OAuth result for tokens.']],
          widths=[1.9, 4.6])

    h2(doc, '6.1 Entity relationships')
    er = r"""
 Category 1---* Institution 1---* Complaint *---1 User (nullable: anonymous)
      |                                |  \
      *                               1*   *---* Attachment
   Complaint                     ComplaintEvent
                                       |
 User 1---* RefreshToken               *
 User 1---* Consent            EmailOutbox *---1 Complaint (related)
 User 1---* OAuthIdentity      AuditLog   *---1 User (actor, nullable)
"""
    mono_block(doc, er.strip('\n'))

    h2(doc, '6.2 Enumerations')
    table(doc, ['Enum', 'Values'],
          [['UserRole', 'admin, user'],
           ['UserStatus', 'draft, pending_confirmation, active, blocked, deactivated'],
           ['ComplaintStatus', 'submitted, pending_review, approved, rejected, forwarded, closed'],
           ['SubmissionType', 'authenticated, anonymous'],
           ['EmailOutboxStatus', 'pending, sent, failed, dead'],
           ['ComplaintEventType', 'created, pending_review, approved, rejected, forwarded, closed, '
            'email_dispatched, email_failed, attachment_added']],
          widths=[1.8, 4.7])

    # 7. API reference
    h1(doc, '7. API Reference')
    para(doc, 'All endpoints are under /api. Authentication uses a Bearer access token; the '
              'refresh token travels in an HttpOnly cookie. Admin routes require the admin role.')
    h2(doc, '7.1 Public & authentication')
    table(doc, ['Method', 'Path', 'Description'],
          [['GET', '/api/health', 'Liveness probe.'],
           ['GET', '/api/config/public', 'Client runtime config (limits, captcha, providers, locales).'],
           ['POST', '/api/auth/register', 'Register (captcha, consent).'],
           ['GET', '/api/auth/verify-email', 'Consume email-verification token.'],
           ['POST', '/api/auth/resend-verification', 'Resend verification email.'],
           ['POST', '/api/auth/login', 'Login; returns access token + refresh cookie.'],
           ['POST', '/api/auth/refresh', 'Rotate refresh token; issue new access token.'],
           ['POST', '/api/auth/logout', 'Revoke refresh token.'],
           ['POST', '/api/auth/forgot-password', 'Request password reset (captcha).'],
           ['POST', '/api/auth/reset-password', 'Set new password; revoke sessions.'],
           ['GET', '/api/auth/oauth/google[/callback]', 'Google OAuth start / callback.'],
           ['GET', '/api/auth/oauth/facebook[/callback]', 'Facebook OAuth start / callback.'],
           ['POST', '/api/auth/oauth/exchange', 'Exchange one-time code for tokens.']],
          widths=[0.7, 2.9, 2.9])
    h2(doc, '7.2 Complaints, catalogue & files')
    table(doc, ['Method', 'Path', 'Description'],
          [['POST', '/api/complaints', 'Submit a complaint (auth optional; multipart).'],
           ['GET', '/api/complaints', 'List own complaints (auth).'],
           ['GET', '/api/complaints/:publicId', 'Complaint detail (owner or admin).'],
           ['GET', '/api/categories', 'List complaint categories.'],
           ['GET', '/api/institutions', 'List/search institutions (paged).'],
           ['GET', '/api/attachments/:id', 'Stream an attachment (owner or admin).'],
           ['GET', '/api/statistics/summary', 'Public summary statistics.']],
          widths=[0.7, 2.9, 2.9])
    h2(doc, '7.3 Administration (role: admin)')
    table(doc, ['Method', 'Path', 'Description'],
          [['GET', '/api/admin/complaints/search', 'Search/filter all complaints.'],
           ['POST', '/api/admin/complaints/:publicId/approve', 'Approve pending complaint.'],
           ['POST', '/api/admin/complaints/:publicId/reject', 'Reject with reason.'],
           ['PATCH', '/api/admin/complaints/:publicId/status', 'Change status (e.g. close).'],
           ['GET', '/api/admin/complaints/:publicId/events', 'Complaint event timeline.'],
           ['POST/PATCH/DELETE', '/api/admin/institutions[/:id]', 'Manage institutions.'],
           ['GET', '/api/admin/users', 'List users.'],
           ['PATCH', '/api/admin/users/:id', 'Update role/status.'],
           ['POST', '/api/admin/users/:id/anonymize', 'Anonymise (GDPR).'],
           ['GET', '/api/admin/users/:id/export', 'Export user data.'],
           ['GET', '/api/admin/statistics/detail', 'Detailed statistics (date range).'],
           ['GET', '/api/admin/statistics/export', 'Export CSV or PDF.']],
          widths=[1.3, 3.0, 2.2])

    # 8. Security architecture
    h1(doc, '8. Security Architecture')
    h2(doc, '8.1 Authentication & sessions')
    kv_bullet(doc, 'Password hashing', 'Argon2id with configurable memory/time/parallelism.')
    kv_bullet(doc, 'Access token', 'Short-lived JWT (default 15m) held in memory by the SPA.')
    kv_bullet(doc, 'Refresh token', 'Long-lived (default 7d), stored hashed (SHA-256) in the DB '
              'and delivered as an HttpOnly, SameSite=Lax cookie.')
    kv_bullet(doc, 'Rotation & reuse detection', 'Every refresh rotates the token; presenting an '
              'already-revoked token revokes the entire token family and raises an audit event.')
    kv_bullet(doc, 'Lockout', 'Five failed logins lock the account for 15 minutes.')
    kv_bullet(doc, 'OAuth safety', 'A state cookie protects the callback; the browser never sees '
              'raw tokens — it exchanges a one-time, 60-second code instead.')
    h2(doc, '8.2 Platform hardening')
    bullet(doc, 'Helmet Content-Security-Policy with frame-ancestors none and object-src none.')
    bullet(doc, 'Credentialed CORS restricted to the configured frontend origin.')
    bullet(doc, 'Global, auth, and complaint-submission rate limiters.')
    bullet(doc, 'hCaptcha on registration and password-reset requests.')
    bullet(doc, 'Attachment uploads validated by real content (magic bytes), not just MIME header.')
    bullet(doc, 'Attachment downloads restricted to the owner or an administrator, and audited.')
    bullet(doc, 'Structured logging with redaction of Authorization and Cookie headers.')

    # 9. Communication flows
    h1(doc, '9. Communication Flows & Sequence Diagrams')

    h2(doc, '9.1 Registration & email verification')
    seq1 = r"""
User        SPA            API/auth          DB            Outbox/Worker   Mailbox
 |  fill form |               |               |                 |            |
 |----------->| POST register |               |                 |            |
 |            |-------------->| verify captcha|                 |            |
 |            |               | create user   |                 |            |
 |            |               |   (pending)   |-- insert ------>|            |
 |            |               | store consents|-- insert ------>|            |
 |            |               | enqueue verify email ---------->| (pending)  |
 |            |  201 created  |               |                 |            |
 |            |<--------------|               |                 |            |
 |            |               |               |  worker polls ->| send mail ->|
 |  click verify link (token) |               |                 |            |
 |--------------------------->| verify-email  |                 |            |
 |            |               | consume token; user -> active   |            |
 |            |  redirect OK  |               |                 |            |
"""
    mono_block(doc, seq1.strip('\n'))

    h2(doc, '9.2 Login with refresh-token rotation')
    seq2 = r"""
SPA                 API/auth                        DB
 | POST /login (email,pw) |                          |
 |----------------------->| find user; check lockout |
 |                        | verify Argon2id password |
 |                        | issue access JWT (15m)   |
 |                        | create refresh (hashed) --> insert
 |                        | reset failedLoginCount   |
 | 200 {accessToken,user} |                          |
 |<-----------------------| Set-Cookie: refresh (HttpOnly)
 |                        |                          |
 | ... access token expires ...                      |
 | POST /refresh (cookie) |                          |
 |----------------------->| verify + look up hash    |
 |                        | if revoked -> revoke ALL (reuse detected)
 |                        | else rotate: revoke old, issue new
 | 200 {accessToken}      | Set-Cookie: new refresh  |
 |<-----------------------|                          |
"""
    mono_block(doc, seq2.strip('\n'))

    h2(doc, '9.3 OAuth sign-in (Google / Facebook)')
    seq3 = r"""
Browser        API/auth           Provider         DB           SPA
 | GET /oauth/google |               |              |            |
 |------------------>| set state cookie; redirect -->|            |
 |                   |               | user consents|            |
 | callback + state  |<--------------|              |            |
 |------------------>| verify state; exchange code  |            |
 |                   | find/create user + identity ->| upsert    |
 |                   | issue one-time exchange code  |            |
 |                   | redirect to SPA /auth/callback?code=...    |
 |------------------------------------------------->|            |
 |                   |               |              |  POST /oauth/exchange
 |                   |<---------------------------------------- (code)
 |                   | consume code; issue tokens   |            |
 |                   | 200 {accessToken} + refresh cookie ------>|
"""
    mono_block(doc, seq3.strip('\n'))

    h2(doc, '9.4 Registered complaint submission & forwarding')
    seq4 = r"""
User    SPA           API/complaints        DB              Outbox        Worker/SMTP
 | fill  |                |                  |                |             |
 |------>| POST complaint |                  |                |             |
 |       |--------------->| optionalAuth (user present)       |             |
 |       |                | validate + sniff attachments      |             |
 |       |                | create complaint (submitted) ---> insert        |
 |       |                | event: created                    |             |
 |       |                | forwardComplaint():               |             |
 |       |                |   TX claim status -> forwarded     |             |
 |       |                |   enqueue -> institution email --> (pending)     |
 |       |                |   enqueue -> OMBUDSMAN email ----> (pending)     |
 |       |                |   enqueue -> submitter copy (opt)  |             |
 |       |                |   events: email_dispatched, forwarded            |
 |       | 201 {publicId} |                  |                |             |
 |<------|                |                  |  poll (15s) -->| send x3 ---->|
"""
    mono_block(doc, seq4.strip('\n'))

    h2(doc, '9.5 Anonymous complaint moderation')
    seq5 = r"""
Anon   SPA          API/complaints     DB           Admin        API/admin
 | fill |               |              |             |              |
 |----->| POST complaint |             |             |              |
 |      |-------------->| no user -> pending_review  |              |
 |      |               | create complaint -> insert |              |
 |      |               | enqueue admin notice email |              |
 |      | 201 {pending} |              |             |              |
 |      |               |              |  review queue|             |
 |      |               |              |<-------------| GET search  |
 |      |               |              |              | approve/reject
 |      |               |              |<------------------------- |
 |      |               | approve: status->approved; forwardComplaint()
 |      |               | reject: status->rejected (reason recorded)|
"""
    mono_block(doc, seq5.strip('\n'))

    h2(doc, '9.6 Asynchronous email delivery (outbox worker)')
    para(doc, 'Emails are never sent inline with the request. They are written to the '
              'email_outbox table inside the same transaction as the complaint change, then '
              'delivered by a background worker. This guarantees the request succeeds or fails '
              'atomically and that delivery is retried on transient SMTP errors.')
    seq6 = r"""
every 15s:
  worker -> pg_try_advisory_lock(key)        # only one instance sends at a time
  if lock acquired:
     fetch up to 20 rows WHERE status=pending AND nextAttemptAt<=now
     for each email:
        try  send via SMTP  -> mark 'sent'
        catch -> attempts++, record error, backoff nextAttemptAt
                 if attempts >= 5 -> mark 'dead' (dead-letter)
     pg_advisory_unlock(key)
"""
    mono_block(doc, seq6.strip('\n'))
    para(doc, 'Email templates: auth.verify_email, auth.password_reset, complaint.to_institution, '
              'complaint.to_ombudsman, complaint.to_user_copy, admin.new_anonymous_complaint.')

    # 10. Configuration
    h1(doc, '10. Configuration Reference')
    para(doc, 'Configuration is supplied through environment variables and validated at startup. '
              'Selected variables (defaults in parentheses):')
    table(doc, ['Variable', 'Purpose (default)'],
          [['NODE_ENV / PORT', 'Runtime mode / API port (development / 3000).'],
           ['DATABASE_URL', 'PostgreSQL connection string.'],
           ['PUBLIC_FRONTEND_URL / PUBLIC_BACKEND_URL', 'Public origins for links and CORS.'],
           ['JWT_ACCESS_SECRET / JWT_REFRESH_SECRET', 'Signing secrets (must be set in prod).'],
           ['JWT_ACCESS_TTL / JWT_REFRESH_TTL', 'Token lifetimes (15m / 7d).'],
           ['ARGON2_MEMORY/TIME/PARALLELISM', 'Password-hash cost (19456 / 2 / 1).'],
           ['GOOGLE_* / FACEBOOK_*', 'OAuth client credentials and redirect URLs.'],
           ['HCAPTCHA_SITE_KEY / HCAPTCHA_SECRET', 'Captcha keys.'],
           ['SMTP_HOST/PORT/USER/PASSWORD/SECURE', 'Mail transport (mailhog / 1025).'],
           ['MAIL_FROM / OMBUDSMAN_EMAIL', 'Sender and mandatory oversight recipient.'],
           ['ADMIN_EMAIL / ADMIN_INITIAL_PASSWORD', 'Seeded administrator account.'],
           ['UPLOAD_DIR / MAX_ATTACHMENTS / MAX_ATTACHMENT_*_BYTES', 'Upload storage and limits (3 / 5 MB).'],
           ['RATE_LIMIT_GLOBAL/AUTH/COMPLAINT_SUBMIT', 'Per-minute limits (300 / 10 / 10).'],
           ['MIN_BODY_LENGTH', 'Minimum complaint description length (100).']],
          widths=[2.7, 3.8])

    # 11. Deployment
    h1(doc, '11. Build, Run & Deployment')
    h2(doc, '11.1 Mode A — everything in Docker')
    mono_block(doc, 'docker compose up --build\n'
                    'docker compose exec backend npm run seed\n\n'
                    'App:      http://localhost:5174\n'
                    'API:      http://localhost:3000/api/health\n'
                    'MailHog:  http://localhost:8025\n'
                    'Postgres: localhost:5433')
    h2(doc, '11.2 Mode B — infra in Docker, app on host (development)')
    mono_block(doc, '# override for host-side runs\n'
                    'DATABASE_URL=postgresql://vallentin:vallentin@localhost:5433/vallentin?schema=public\n'
                    'SMTP_HOST=localhost\n'
                    'UPLOAD_DIR=./uploads\n\n'
                    'docker compose up -d postgres mailhog\n'
                    'npm install\n'
                    'cd backend && npx prisma migrate deploy && npm run seed\n'
                    'cd backend && npm run dev      # terminal 1 (API, watch)\n'
                    'cd frontend && npm run dev     # terminal 2 (Vite HMR)')
    h2(doc, '11.3 Backend & frontend scripts')
    table(doc, ['Command', 'Effect'],
          [['backend: npm run dev', 'Start API with tsx watch (hot reload).'],
           ['backend: npm run build', 'Compile TypeScript to dist/.'],
           ['backend: npm start', 'Run compiled server (node dist/server.js).'],
           ['backend: npm run prisma:migrate', 'Apply/create Prisma migrations.'],
           ['backend: npm run seed', 'Seed categories, admin, demo data.'],
           ['backend: npm test', 'Jest + Supertest suite.'],
           ['frontend: npm run dev', 'Vite dev server on :5174.'],
           ['frontend: npm run build', 'Type-check then production build.'],
           ['frontend: npm run test / test:e2e', 'Unit tests / Playwright E2E.']],
          widths=[2.6, 3.9])

    # 12. Testing
    h1(doc, '12. Testing Strategy')
    bullet(doc, 'Backend integration/unit tests (Jest + Supertest) cover auth, complaints, admin, '
                'statistics, and the email worker, plus a smoke test.')
    bullet(doc, 'Frontend unit tests run under Vitest; end-to-end journeys use Playwright.')
    bullet(doc, 'Live-backend E2E specs are gated behind the E2E_LIVE_BACKEND flag.')
    bullet(doc, 'Type-checking (tsc --noEmit) and ESLint run per workspace.')

    # 13. Observability & ops
    h1(doc, '13. Observability & Operations')
    kv_bullet(doc, 'Logging', 'Structured JSON logs via Pino with per-request IDs; 4xx as warn, '
              '5xx as error; Authorization and Cookie headers redacted.')
    kv_bullet(doc, 'Health', 'GET /api/health for liveness checks.')
    kv_bullet(doc, 'Audit trail', 'Security and lifecycle events persisted to audit_log.')
    kv_bullet(doc, 'Graceful shutdown', 'SIGTERM/SIGINT stop the worker, drain the server, and '
              'disconnect Prisma, with a forced-exit safety timeout.')
    kv_bullet(doc, 'Email resilience', 'Outbox with retry and dead-letter; single-sender guarantee '
              'via a PostgreSQL advisory lock.')

    h2(doc, '13.1 Troubleshooting quick reference')
    table(doc, ['Symptom', 'Likely cause / fix'],
          [['Login returns AUTH_INVALID_CREDENTIALS', 'Seed not run — run npm run seed in backend.'],
           ['Port 5432 in use', 'Postgres is mapped to 5433 by design; adjust if also taken.'],
           ['Cannot find module @prisma/client', 'Run npx prisma generate in backend.'],
           ['EPERM mkdir on uploads (Windows)', 'Set UPLOAD_DIR=./uploads for host-side runs.'],
           ['Invalid environment configuration', 'Ensure .env values are exported into the shell.']],
          widths=[2.7, 3.8])

    # 14. Repo layout
    h1(doc, '14. Repository Layout')
    mono_block(doc,
               'backend/    Express API, Prisma schema & migrations, business logic, tests\n'
               '  src/\n'
               '    config/          env validation\n'
               '    lib/             jwt, password, ids, captcha, csv, pdf, logger\n'
               '    middleware/      auth, rbac, validate, captcha, rateLimit, errorHandler\n'
               '    modules/         auth, users, complaints, institutions, statistics,\n'
               '                     attachments, email, ombudsman, consents, auditLog, admin\n'
               '    app.ts / server.ts\n'
               '  prisma/            schema.prisma, migrations, seed\n'
               'frontend/   React SPA (Vite)\n'
               '  src/\n'
               '    pages/ components/ api/ stores/ hooks/ utils/ i18n\n'
               'docs/       ARCHITECTURE.md, API_CONTRACT.md, DATA_MODEL.md')

    path = os.path.join(OUT_DIR, 'Vallentin_Claims_Technical_Documentation.docx')
    doc.save(path)
    return path


if __name__ == '__main__':
    b = build_business()
    print('Wrote', b)
    t = build_technical()
    print('Wrote', t)
